from datetime import datetime

import pandas as pd
import pytest
import pytz
from openpyxl import Workbook, load_workbook

import pool


PACIFIC = pytz.timezone("America/Los_Angeles")


def schedule_frame(spreads=None):
    spreads = spreads or [2.5, 3.5, 7.5]
    frame = pd.DataFrame(
        {
            "Team1": ["SEAHAWKS", "RAMS", "JAGUARS"],
            "Spread": [-spreads[0], -spreads[1], -spreads[2]],
            "Team2": ["PATRIOTS", "49ERS", "BROWNS"],
            "Team1_Abbr": ["SEA", "LAR", "JAC"],
            "Team2_Abbr": ["NE", "SF", "CLE"],
            "Home_Team": ["PATRIOTS", "49ERS", "BROWNS"],
            "UTC_DateTime": [
                "2026-09-10T00:20:00Z",  # Wednesday Pacific
                "2026-09-11T00:20:00Z",  # Thursday Pacific
                "2026-09-13T17:00:00Z",  # Sunday Pacific
            ],
            "Favorite_Side": ["away", "away", "away"],
        }
    )
    frame["game_day"] = (
        pd.to_datetime(frame["UTC_DateTime"], utc=True)
        .dt.tz_convert("America/Los_Angeles")
        .dt.day_name()
    )
    return frame


def create_week_workbook(path):
    workbook = Workbook()
    template = workbook.active
    template.title = "Template"
    week = workbook.copy_worksheet(template)
    week.title = "Week 1"
    for row, favorite, spread, underdog, favorite_abbr, underdog_abbr in (
        (2, "SEAHAWKS", 2.5, "PATRIOTS", "SEA", "NE"),
        (3, "RAMS", 3.5, "49ERS", "LAR", "SF"),
        (4, "JAGUARS", 7.5, "BROWNS", "JAC", "CLE"),
    ):
        week.cell(row=row, column=3).value = favorite
        week.cell(row=row, column=4).value = spread
        week.cell(row=row, column=5).value = underdog
        week.cell(row=row, column=9).value = favorite_abbr
        week.cell(row=row, column=11).value = underdog_abbr
        week.cell(row=row, column=14).value = f"HM-{row}"
        week.cell(row=row, column=15).value = f"DE-{row}"
    workbook.save(path)


def matchup_cells(path, row):
    workbook = load_workbook(path, read_only=True, data_only=False)
    sheet = workbook["Week 1"]
    values = tuple(sheet.cell(row=row, column=column).value for column in (3, 4, 5, 9, 11))
    workbook.close()
    return values


def entry_cells(path, row):
    workbook = load_workbook(path, read_only=True, data_only=False)
    sheet = workbook["Week 1"]
    values = tuple(sheet.cell(row=row, column=column).value for column in (14, 15))
    workbook.close()
    return values


def test_existing_matchups_keep_rows_when_source_order_changes():
    workbook = Workbook()
    sheet = workbook.active
    for row, team1, team2 in ((2, "SEA", "NE"), (3, "LAR", "SF"), (4, "JAC", "CLE")):
        sheet.cell(row=row, column=9).value = team1
        sheet.cell(row=row, column=11).value = team2

    reordered = schedule_frame().iloc[[2, 0, 1]].copy()
    rows = pool.align_excel_rows_to_worksheet(reordered, sheet)

    assert rows.tolist() == [4, 2, 3]


def test_duplicate_current_matchup_aborts_before_row_reuse():
    sheet = Workbook().active
    duplicated = pd.concat([schedule_frame().iloc[[0]], schedule_frame().iloc[[0]]])

    with pytest.raises(ValueError, match="duplicate matchups.*NE vs SEA"):
        pool.align_excel_rows_to_worksheet(duplicated, sheet)


def test_duplicate_existing_matchup_aborts_before_workbook_update():
    sheet = Workbook().active
    for row in (2, 3):
        sheet.cell(row=row, column=9).value = "SEA"
        sheet.cell(row=row, column=11).value = "NE"

    with pytest.raises(ValueError, match="duplicate matchup NE vs SEA.*rows 2 and 3"):
        pool.align_excel_rows_to_worksheet(schedule_frame().iloc[[0]], sheet)


def test_thursday_run_preserves_wednesday_and_updates_future_thursday_game(
    tmp_path, monkeypatch
):
    workbook_path = tmp_path / "Family Football Pool 2026.xlsx"
    create_week_workbook(workbook_path)
    monkeypatch.setenv("file_path", str(workbook_path))
    monkeypatch.setattr(pool, "send_error_email", lambda **kwargs: None)

    updated = schedule_frame(spreads=[4.0, 5.0, 9.0])
    remaining, dotw = pool.prepare_schedule_for_excel(
        updated,
        now=PACIFIC.localize(datetime(2026, 9, 10, 9, 0)),
    )
    pool.update_excel("Week 1", remaining, dotw)

    assert remaining["Team1_Abbr"].tolist() == ["LAR", "JAC"]
    assert matchup_cells(workbook_path, 2) == ("SEAHAWKS", 2.5, "PATRIOTS", "SEA", "NE")
    assert matchup_cells(workbook_path, 3) == ("RAMS", 5, "49ERS", "LAR", "SF")
    assert matchup_cells(workbook_path, 4) == ("JAGUARS", 9, "BROWNS", "JAC", "CLE")
    assert entry_cells(workbook_path, 3) == ("HM-3", "DE-3")
    assert entry_cells(workbook_path, 4) == ("HM-4", "DE-4")


@pytest.mark.parametrize(
    "now",
    [
        PACIFIC.localize(datetime(2026, 9, 11, 9, 0)),
        PACIFIC.localize(datetime(2026, 9, 12, 9, 0)),
    ],
    ids=["friday", "saturday"],
)
def test_friday_and_saturday_runs_preserve_wednesday_and_thursday_rows(
    tmp_path, monkeypatch, now
):
    workbook_path = tmp_path / "Family Football Pool 2026.xlsx"
    create_week_workbook(workbook_path)
    monkeypatch.setenv("file_path", str(workbook_path))
    monkeypatch.setattr(pool, "send_error_email", lambda **kwargs: None)

    updated = schedule_frame(spreads=[4.0, 5.0, 9.0])
    remaining, dotw = pool.prepare_schedule_for_excel(updated, now=now)
    pool.update_excel("Week 1", remaining, dotw)

    assert remaining["Team1_Abbr"].tolist() == ["JAC"]
    assert matchup_cells(workbook_path, 2) == ("SEAHAWKS", 2.5, "PATRIOTS", "SEA", "NE")
    assert matchup_cells(workbook_path, 3) == ("RAMS", 3.5, "49ERS", "LAR", "SF")
    assert matchup_cells(workbook_path, 4) == ("JAGUARS", 9, "BROWNS", "JAC", "CLE")
    assert entry_cells(workbook_path, 4) == ("HM-4", "DE-4")
